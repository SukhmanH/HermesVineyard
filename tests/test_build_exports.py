"""Monthly audit export tests: real schema, isolated DB, explicit historical period."""
from __future__ import annotations

import importlib.util
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent


@pytest.fixture
def engine(db, monkeypatch):
    spec = importlib.util.spec_from_file_location("build_exports", REPO / "tools/build_exports.py")
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    db.commit()
    path = db.execute("PRAGMA database_list").fetchone()[2]
    monkeypatch.setattr(mod, "DB_PATH", path)
    return mod


def spray(db, day, corrects=None):
    block = db.execute("SELECT id FROM blocks WHERE code='B3'").fetchone()[0]
    contact = db.execute("SELECT id FROM contacts ORDER BY id LIMIT 1").fetchone()[0]
    return db.execute(
        "INSERT INTO spray_log(log_date,start_time,block_id,acres_treated,product_name_raw,"
        "applicator_contact_id,applicator_name,raw_message,corrects_log_id) "
        "VALUES (?, '08:00', ?, 1, 'Label', ?, 'Worker', '=1+1', ?)",
        (day, block, contact, corrects),
    ).lastrowid


def rows(ws):
    values = list(ws.values)
    return [dict(zip(values[0], r, strict=True)) for r in values[1:]]


def test_monthly_preserves_raw_records_and_cross_month_corrections(engine, db, tmp_path):
    first = spray(db, '2025-09-02')
    second = spray(db, '2025-09-03', first)
    third = spray(db, '2025-10-01', second)
    db.commit()
    out = tmp_path / 'compliance.xlsx'
    assert engine.build_compliance_month(out, 2025, 9) == 2
    wb = openpyxl.load_workbook(out)
    exported = rows(wb['spray_log'])
    by_id = {r['id']: r for r in exported}
    assert set(by_id) == {first, second}
    assert by_id[first]['SUPERSEDED-BY'] == str(second)
    assert by_id[second]['SUPERSEDED-BY'] == str(third)
    columns = {r[1] for r in db.execute('PRAGMA table_info(spray_log)')}
    assert columns <= set(by_id[first])
    raw_col = list(next(wb['spray_log'].values)).index('raw_message') + 1
    assert wb['spray_log'].cell(2, raw_col).data_type == 's'
    assert wb['spray_log'].cell(2, raw_col).value == '=1+1'
    wb.close()


def test_audit_month_uses_vancouver_midnight(engine, db, tmp_path):
    for when in ['2025-09-01T06:59:59Z', '2025-09-01T07:00:00Z',
                 '2025-10-01T06:59:59Z', '2025-10-01T07:00:00Z']:
        db.execute("INSERT INTO audit_log(at_utc,actor,action) VALUES (?, 'test', 'boundary')", (when,))
    db.commit()
    out = tmp_path / 'compliance.xlsx'
    engine.build_compliance_month(out, 2025, 9)
    wb = openpyxl.load_workbook(out)
    assert [r['at_utc'] for r in rows(wb['audit_log'])] == [
        '2025-09-01T07:00:00Z', '2025-10-01T06:59:59Z']
    assert dict(wb['metadata'].values)['period'] == '2025-09'
    wb.close()


def test_invalid_month_does_not_create_file(engine, tmp_path):
    out = tmp_path / 'invalid.xlsx'
    with pytest.raises(ValueError):
        engine.build_compliance_month(out, 2025, 13)
    assert not out.exists()


def test_cli_monthly_export_does_not_run_backups(engine, tmp_path):
    outdir = tmp_path / 'exports'
    env = {**os.environ, 'DB_PATH': engine.DB_PATH, 'EXPORT_DIR': str(outdir),
           'BACKUP_DIR': str(tmp_path / 'backups')}
    run = subprocess.run([sys.executable, str(REPO / 'tools/build_exports.py'),
                          '--compliance-month', '2025-09'],
                         env=env, capture_output=True, text=True, timeout=30)
    assert run.returncode == 0, run.stdout + run.stderr
    assert (outdir / 'compliance-2025-09.xlsx').is_file()
    assert not (tmp_path / 'backups').exists()
